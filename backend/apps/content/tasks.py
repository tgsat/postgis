import json
import os
import re
import shutil
import subprocess
import zipfile

from celery import shared_task
from django.conf import settings
from django.contrib.gis.geos import GEOSGeometry
from django.db import connection, transaction

from .models import ImportJob, Item, default_item_name, normalize_geom_type, unique_title

logger = __import__("logging").getLogger(__name__)

ALLOWED_FIELD_COLS = {"id", "fid", "ogc_fid", "geom", "geojson", "wkb_geometry"}
LON_ALIAS = {"lon", "lng", "long", "longitude", "x", "easting", "east", "xcoord", "ex", "e2"}
LAT_ALIAS = {"lat", "latitude", "y", "northing", "north", "ycoord", "ny", "n2"}
MAX_ZIP_UNCOMPRESSED = 2 * 1024 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000


def _pg_conn():
    return (
        f"PG:host={os.environ.get('POSTGRES_HOST', 'postgis')} "
        f"dbname={os.environ.get('POSTGRES_DB', 'gisdb')} "
        f"user={os.environ.get('POSTGRES_USER', 'postgres')} "
        f"password={os.environ.get('POSTGRES_PASSWORD', 'postgres')} "
        f"port={os.environ.get('POSTGRES_PORT', '5432')}"
    )


def _run(cmd):
    proc = subprocess.run(cmd, capture_output=True, text=True)
    return proc.returncode, proc.stdout, proc.stderr


def _staging_dir(job):
    path = settings.STORAGE_ROOT / "uploads" / str(job.id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _safe_extract_zip(src_path, dest_dir):
    with zipfile.ZipFile(src_path) as zf:
        infos = zf.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise ValueError("ZIP terlalu banyak isi file (maksimum 5000).")
        total = sum(i.file_size for i in infos)
        if total > MAX_ZIP_UNCOMPRESSED:
            raise ValueError("ZIP terlalu besar setelah diekstrak (maksimum 2 GB).")
        for info in infos:
            name = info.filename
            if name.startswith(("/", "\\")) or ".." in name.split("/"):
                raise ValueError(f"Path tidak aman di dalam ZIP: {name!r}")
        zf.extractall(dest_dir)


def _locate_payload(dest_dir):
    """Find the actual datasource inside a staging dir. Returns (kind, path)."""
    shp_files = []
    gdb_dirs = []
    for root, dirs, files in os.walk(dest_dir):
        for f in files:
            low = f.lower()
            if low.endswith(".shp"):
                shp_files.append(os.path.join(root, f))
            elif low.endswith(".gpkg"):
                return "gpkg", os.path.join(root, f)
            elif low.endswith(".csv"):
                return "csv", os.path.join(root, f)
            elif low.endswith(".xlsx"):
                return "xlsx", os.path.join(root, f)
        for d in list(dirs):
            if d.lower().endswith(".gdb"):
                gdb_dirs.append(os.path.join(root, d))
    if shp_files:
        return "shp", shp_files[0]
    if gdb_dirs:
        return "gdb", gdb_dirs[0]
    for root, dirs, files in os.walk(dest_dir):
        for f in files:
            if f.lower().endswith(".gdbtable"):
                return "gdb", root
    raise ValueError("Tidak ditemukan file spasial yang didukung di dalam ZIP.")


def _ogrinfo_layers(datasource):
    rc, out, err = _run(["ogrinfo", datasource])
    layers = []
    if rc != 0:
        raise ValueError(err.strip() or out)
    for line in out.splitlines():
        line = line.strip()
        m = re.match(r"^(\d+):\s+(.+?)\s+\(([^)]+)\)$", line)
        if m:
            layers.append(m.group(2))
            continue
        m2 = re.match(r"^(\d+):\s+(.+)$", line)
        if m2:
            layers.append(m2.group(2))
    if not layers:
        raise ValueError(f"Tidak ada layer terbaca pada datasource: {datasource}")
    return layers


def _ogrinfo_meta(datasource, layer):
    rc, out, err = _run(["ogrinfo", "-so", "-geom=NO", datasource, layer])
    if rc != 0:
        raise ValueError(err.strip() or out)
    meta = {"name": layer, "geom_type": "Geometry", "crs": None, "feature_count": 0, "fields": []}
    geom_m = re.search(r"^Geometry:\s*(.+)$", out, re.MULTILINE)
    if geom_m:
        meta["geom_type"] = normalize_geom_type(geom_m.group(1))
    count_m = re.search(r"^Feature Count:\s*(\d+)$", out, re.MULTILINE)
    if count_m:
        meta["feature_count"] = int(count_m.group(1))
    auth = re.findall(r'(?:AUTHORITY\["EPSG","(\d+)"\]|ID\["EPSG",(\d+)\])', out)
    if auth:
        meta["crs"] = int(auth[-1][0] or auth[-1][1])
    for line in out.splitlines():
        line = line.strip()
        m = re.match(r"^([A-Za-z_][\w]*):\s+(String|Integer|Integer64|Real|Date|DateTime|Time)\s*(?:\(([\d.,]+)\))?$", line)
        if m:
            fname, ftype, size = m.group(1), m.group(2), m.group(3)
            ftype_map = {
                "String": "text",
                "Integer": "number",
                "Integer64": "number",
                "Real": "number",
                "Date": "date",
                "DateTime": "date",
                "Time": "text",
            }
            meta["fields"].append(
                {
                    "name": fname,
                    "label": fname,
                    "type": ftype_map.get(ftype, "text"),
                    "length": int(float(size.split(",")[0])) if size else None,
                    "default": None,
                    "required": False,
                    "domain": [],
                }
            )
    return meta


def _detect_xy_columns(headers):
    lon_col = lat_col = None
    for header in headers:
        key = re.sub(r"[^a-z0-9]", "", str(header).lower())
        if key in LON_ALIAS and lon_col is None:
            lon_col = header
        if key in LAT_ALIAS and lat_col is None:
            lat_col = header
    return lon_col, lat_col


def _read_tabular(path, container):
    """Return (headers, rows_sample_all) for CSV/XLSX."""
    if container == "csv":
        import csv

        with open(path, newline="", encoding="utf-8", errors="replace") as fh:
            reader = csv.reader(fh)
            headers = next(reader, [])
            rows = [row for row in reader]
        return headers, rows
    if container == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers_row = next(it, ())
        headers = [str(h) if h is not None else "" for h in headers_row]
        rows = [list(r) for r in it]
        wb.close()
        return headers, rows
    raise ValueError("Unknown tabular container")


def _tabular_meta(headers, rows):
    lon_col, lat_col = _detect_xy_columns(headers)
    fields = []
    for idx, header in enumerate(headers):
        sample_type = "text"
        for row in rows[:50]:
            if idx < len(row) and row[idx] not in (None, ""):
                val = row[idx]
                try:
                    float(str(val).replace(",", "."))
                    sample_type = "number"
                except (TypeError, ValueError):
                    if isinstance(val, str) and re.match(r"^\d{4}-\d{2}-\d{2}", val.strip()):
                        sample_type = "date"
                    else:
                        sample_type = "text"
                break
        fields.append(
            {
                "name": re.sub(r"[^A-Za-z0-9]+", "_", (header or f"field_{idx}")).strip("_").lower() or f"field_{idx}",
                "label": header,
                "type": sample_type,
                "length": None,
                "default": None,
                "required": False,
                "domain": [],
            }
        )
    geom = "Point" if lon_col and lat_col else "Geometry"
    return {
        "name": "tabular",
        "geom_type": geom,
        "crs": 4326 if geom == "Point" else None,
        "feature_count": len(rows),
        "fields": fields,
        "lon_col": lon_col,
        "lat_col": lat_col,
        "needs_coordinates": not (lon_col and lat_col),
    }


def _ogr2ogr_staging(job, datasource, layer, table):
    rc, out, err = _run(
        [
            "ogr2ogr",
            "-f", "PostgreSQL",
            _pg_conn(),
            datasource,
            layer,
            "-nln", table,
            "-t_srs", "EPSG:4326",
            "-lco", "GEOMETRY_NAME=geom",
            "-overwrite",
            "-lco", "SCHEMA=public",
        ]
    )
    if rc != 0:
        tail = (err or out).strip().splitlines()
        raise ValueError("ogr2ogr gagal: " + " | ".join(tail[:5]))


def _read_staging_features(table, field_mapping=None):
    """Yield (geom, props) tuples read from a staging PostGIS table."""
    field_mapping = field_mapping or {}
    with connection.cursor() as cursor:
        cursor.execute(f'SELECT ST_AsBinary(geom) AS wkb, to_jsonb(t) - \'geom\' AS props FROM "{table}" t')
        while True:
            rows = cursor.fetchmany(2000)
            if not rows:
                break
            for wkb, props in rows:
                if not wkb:
                    continue
                geom = GEOSGeometry(bytes(wkb).hex(), srid=4326)
                props = json.loads(props) if isinstance(props, str) else (props or {})
                clean = {k: v for k, v in props.items() if k.lower() not in ALLOWED_FIELD_COLS}
                clean = {field_mapping.get(k, k): v for k, v in clean.items()}
                yield geom, clean
    with connection.cursor() as cursor:
        cursor.execute(f'DROP TABLE IF EXISTS "{table}"')


def _build_features(dataset, features):
    from apps.layers.models import Feature

    user = dataset.created_by
    data = [Feature(dataset=dataset, geom=geom, props=dict(props), created_by=user) for geom, props in features]
    for i in range(0, len(data), 1000):
        Feature.objects.bulk_create(data[i : i + 1000])


def _create_dataset(job, layer_meta, title, project, field_mapping, created_by):
    from apps.accounts.models import User
    from apps.layers.models import Dataset

    base = re.sub(r"[^a-z0-9]+", "_", (title or "layer").lower()).strip("_")[:50]
    table = base or "layer"
    collision = Dataset.objects.filter(table_name__startswith=table).count()
    if collision:
        table = f"{table}_{collision + 1}"

    dataset = Dataset.objects.create(
        company=job.company,
        project=project,
        name=title,
        table_name=table,
        geom_type=layer_meta.get("geom_type", "Geometry"),
        srid=4326,
        schema={"fields": layer_meta.get("fields", []), "source": job.file_name},
        is_public=False,
        is_published=True,
        created_by=created_by,
    )
    return dataset


def _set_job_status(job, status, error=""):
    job.status = status
    job.error = error
    if status in ("success", "failed"):
        from django.utils import timezone

        job.finished_at = timezone.now()
    job.save(update_fields=["status", "error", "finished_at"])


@shared_task
def analyze_import_job(job_id, user_id=None):
    job = ImportJob.objects.get(pk=job_id)
    _set_job_status(job, "analyzing")
    try:
        staging = _staging_dir(job)
        datasource = job.file_path
        container = job.container_kind
        base = job.file_name.rsplit(".", 1)[0]

        if container == "zip":
            _safe_extract_zip(job.file_path, staging)
            kind, payload = _locate_payload(staging)
            container = kind
            datasource = payload

        layers = []
        if container in ("shp", "gdb", "gpkg"):
            layernames = _ogrinfo_layers(datasource)
            single = len(layernames) == 1
            for layername in layernames:
                meta = _ogrinfo_meta(datasource, layername)
                meta["title"] = default_item_name(job.file_name) if single else default_item_name(f"{base}_{layername}")
                layers.append(meta)
        elif container in ("csv", "xlsx"):
            headers, rows = _read_tabular(datasource, container)
            meta = _tabular_meta(headers, rows)
            meta["title"] = default_item_name(job.file_name)
            layers.append(meta)
        else:
            raise ValueError(f"Container tidak dikenali: {container}")

        job.format = container
        job.layers = layers
        job.file_path = datasource
        job.status = "ready"
        job.error = ""
        job.save(update_fields=["format", "layers", "file_path", "status", "error"])
        return {"job": job_id, "layers": len(layers), "status": "ready"}
    except Exception as exc:
        logger.exception("analyze_import_job failed")
        _set_job_status(job, "failed", str(exc))
        return {"job": job_id, "status": "failed", "error": str(exc)}


@shared_task
def publish_import_job(job_id, user_id=None):
    from apps.accounts.models import User
    from apps.projects.models import Project
    from apps.rbac.models import AuditLog

    job = ImportJob.objects.get(pk=job_id)
    _set_job_status(job, "processing")
    user = User.objects.filter(pk=user_id).first() if user_id else None
    created_item_ids = []
    try:
        settings_conf = job.config or {}
        requested = settings_conf.get("layers") or []
        project = Project.objects.filter(pk=settings_conf.get("project")).first() if settings_conf.get("project") else None
        field_mapping = settings_conf.get("field_mapping") or {}

        staging = _staging_dir(job)
        if requested:
            metas = [layer for layer in job.layers if layer.get("name") in {r.get("name") for r in requested}]
            if not metas:
                raise ValueError("Tidak ada layer yang dipilih untuk dipublish.")
        else:
            metas = job.layers

        for index, meta in enumerate(metas):
            title = None
            for req in requested:
                if req.get("name") == meta.get("name"):
                    title = req.get("title") or None
                    break
            if not title:
                title = meta.get("title") or default_item_name(f"{job.file_name}_{meta['name']}")
            title = unique_title(job.company, title)

            dataset = _create_dataset(job, meta, title, project, field_mapping, user)

            if job.format in ("shp", "gdb", "gpkg"):
                table = f"staging_job_{job.id}_{index}"
                src = job.file_path
                layer_name = meta["name"]
                _ogr2ogr_staging(job, src, layer_name, table)
                features = _read_staging_features(table, field_mapping)
                _build_features(dataset, features)
            elif job.format in ("csv", "xlsx"):
                headers, rows = _read_tabular(job.file_path, job.format)
                lon_col = meta.get("lon_col")
                lat_col = meta.get("lat_col")
                for req in requested:
                    if req.get("lon_col"):
                        lon_col = req["lon_col"]
                    if req.get("lat_col"):
                        lat_col = req["lat_col"]
                if not lon_col or not lat_col:
                    raise ValueError("Butuh kolom koordinat (lon/lat) untuk CSV/XLSX. Tidak ditemukan kolom koordinat yang sah.")
                lon_idx = headers.index(lon_col) if lon_col in headers else None
                lat_idx = headers.index(lat_col) if lat_col in headers else None
                if lon_idx is None or lat_idx is None:
                    raise ValueError("Kolom koordinat tidak ditemukan pada baris header.")
                header_lookup = {h: i for i, h in enumerate(headers)}
                feats = []
                for row in rows:
                    props = {}
                    for h, idx in header_lookup.items():
                        if idx < len(row):
                            key = field_mapping.get(h, re.sub(r"[^A-Za-z0-9]+", "_", str(h)).strip("_").lower() or f"field_{idx}")
                            props[key] = row[idx]
                    try:
                        lon = float(str(row[lon_idx]).replace(",", "."))
                        lat = float(str(row[lat_idx]).replace(",", "."))
                    except (TypeError, ValueError, IndexError) as exc:
                        raise ValueError(f"Baris dengan koordinat tidak valid: lon={row[lon_idx] if lon_idx < len(row) else None}, lat={row[lat_idx] if lat_idx < len(row) else None}") from exc
                    geom = GEOSGeometry(f"POINT({lon} {lat})", srid=4326)
                    feats.append((geom, dict(props)))
                _build_features(dataset, feats)
            else:
                raise ValueError(f"Format tidak didukung saat publish: {job.format}")

            item = Item.objects.get(dataset=dataset)
            created_item_ids.append(item.pk)
            AuditLog.objects.create(
                user=user,
                action="IMPORT",
                resource_type="dataset",
                resource_id=str(dataset.pk),
                detail=f'Add Data published dataset "{dataset.name}"',
                ip=None,
            )

        job.item_ids = created_item_ids
        _set_job_status(job, "success")
        job.save(update_fields=["item_ids"])
        shutil.rmtree(staging, ignore_errors=True)
        return {"job": job_id, "items": created_item_ids, "status": "success"}
    except Exception as exc:
        logger.exception("publish_import_job failed")
        _set_job_status(job, "failed", str(exc))
        return {"job": job_id, "status": "failed", "error": str(exc)}